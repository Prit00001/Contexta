"""Builds tiny PDF/DOCX fixtures in-memory (no checked-in binaries, no network) to
exercise each parser."""
from core.parsers import parse_docx, parse_document, parse_pdf, parse_text


def test_parse_text_tracks_markdown_headings():
    content = b"# Intro\nFirst paragraph.\n\n## Details\nSecond paragraph."
    blocks = parse_text(content)
    assert [b.section_heading for b in blocks] == ["Intro", "Details"]
    assert blocks[0].page_number is None


def test_parse_pdf_extracts_text_and_page_numbers():
    import fitz

    doc = fitz.open()
    for text in ("Page one content.", "Page two content."):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()

    blocks = parse_pdf(data)
    assert [b.page_number for b in blocks] == [1, 2]
    assert "Page one" in blocks[0].text


def test_parse_docx_tracks_headings():
    import io

    import docx

    d = docx.Document()
    d.add_heading("Section A", level=1)
    d.add_paragraph("Body text under section A.")
    buf = io.BytesIO()
    d.save(buf)

    blocks = parse_docx(buf.getvalue())
    assert blocks[0].section_heading == "Section A"
    assert "Body text" in blocks[0].text


def test_parse_document_dispatches_by_content_type():
    blocks = parse_document(b"plain text content", "text/plain")
    assert blocks[0].text == "plain text content"
